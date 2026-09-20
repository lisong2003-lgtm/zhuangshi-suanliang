#!/usr/bin/env python3
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

SKILL = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("decor_qty", SKILL / "scripts" / "decor_qty.py")
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

rules = json.loads((SKILL / "rules" / "装饰做法.json").read_text(encoding="utf-8"))
standards = json.loads((SKILL / "rules" / "国家规范要求.json").read_text(encoding="utf-8"))
project = json.loads((SKILL / "templates" / "项目专属规则示例.json").read_text(encoding="utf-8"))
room_rules = json.loads((SKILL / "rules" / "标准房间模板.json").read_text(encoding="utf-8"))
quantity_rules = json.loads((SKILL / "rules" / "清单工程量规则.json").read_text(encoding="utf-8"))
price_rules = json.loads((SKILL / "rules" / "概算综合单价.json").read_text(encoding="utf-8"))
material_rules = json.loads((SKILL / "rules" / "物资分类规则.json").read_text(encoding="utf-8"))


def test_ground_area():
    area, unit, formula, status, missing = m.base_quantity({
        "部位类型": "地面/楼面", "长_m": "6", "宽_m": "4.5", "间数": "2"
    })
    assert area == 54 and unit == "m2" and status == "完成" and not missing


def test_wall_openings():
    area, unit, formula, status, missing = m.base_quantity({
        "部位类型": "墙面", "周长_m": "17.2", "高_m": "3.1",
        "门窗洞口面积_m2": "1.8", "洞口侧壁增加_m2": "0.6", "间数": "1"
    })
    assert abs(area - (17.2 * 3.1 - 1.8 + 0.6)) < 1e-9
    assert unit == "m2" and status == "完成"


def test_material_mapping():
    row = {"部位类型": "地面/楼面", "做法编号": "D01", "长_m": "6", "宽_m": "4.5", "间数": "1"}
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    assert result["基面量"] == 27 and result["状态"] == "完成" and details
    names = {d["材料名称"] for d in details}
    assert {"地砖", "水泥砂浆", "填缝剂"} <= names


def test_missing_finish():
    row = {"部位类型": "地面/楼面", "做法编号": "X99", "长_m": "6", "宽_m": "4.5", "间数": "1"}
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    assert result["状态"] == "缺项" and "做法编号 X99 不在规则表" in result["缺项说明"]


def test_dedicated_missing_node():
    row = {
        "部位类型": "防水", "做法编号": "F01", "做法类型": "项目专属",
        "长_m": "2.4", "宽_m": "1.8", "周长_m": "8.4", "间数": "1"
    }
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    assert result["状态"] == "缺项"
    assert "项目专属做法缺少节点编号或节点详图" in result["缺项说明"]


def test_waterproof_coating_additional_and_thickness():
    row = {
        "部位类型": "防水", "做法编号": "F01", "长_m": "2.4", "宽_m": "1.8",
        "周长_m": "8.4", "防水上翻高度_m": "0.3", "附加层面积_m2": "2.1",
        "涂膜厚度_mm": "1.5", "附加层宽度_m": "0.25", "搭接宽度_mm": "100",
        "卷材幅宽_mm": "1000", "间数": "1"
    }
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    assert abs(result["基面量"] - (2.4 * 1.8 + 8.4 * 0.3)) < 1e-9
    assert result["状态"] == "完成"
    add = next(d for d in details if d["材料名称"] == "防水附加层")
    assert abs(add["订货量"] - 2.1 * 1.05) < 1e-9
    assert "涂膜厚度 1.5mm ≥ 参考下限 1.5mm" in result["防水校验"]


def test_waterproof_membrane_overlap_and_additional():
    row = {
        "部位类型": "屋面防水", "做法编号": "F02", "长_m": "5", "宽_m": "6",
        "周长_m": "22", "防水上翻高度_m": "0.25", "附加层长度_m": "12",
        "附加层宽度_m": "0.5", "搭接宽度_mm": "100", "卷材幅宽_mm": "1000",
        "防水道数": "2", "涂膜厚度_mm": "1.5", "间数": "1"
    }
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    assert result["状态"] == "完成"
    membrane = next(d for d in details if d["材料计算方式"] == "overlap_area")
    assert abs(membrane["订货量"] - (5 * 6 + 22 * 0.25) * (1000 / 900) * 1.1) < 1e-9
    additional = next(d for d in details if d["材料名称"] == "防水附加层")
    assert abs(additional["订货量"] - 12 * 0.5 * 1.05) < 1e-9


def test_waterproof_membrane_lap_below_minimum():
    row = {
        "部位类型": "防水", "做法编号": "F02", "长_m": "5", "宽_m": "6",
        "周长_m": "22", "附加层长度_m": "12", "附加层宽度_m": "0.5",
        "搭接宽度_mm": "70", "卷材幅宽_mm": "1000", "涂膜厚度_mm": "1.5", "间数": "1"
    }
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    assert result["状态"] == "缺项"
    assert "搭接宽度 70mm 小于内置参考下限 80mm" in result["缺项说明"]


def test_project_node_rule_overrides_default():
    row = {
        "部位类型": "防水", "做法编号": "F01", "做法类型": "项目专属", "节点编号": "JC-01",
        "节点详图": "建施-15 卫生间防水节点", "长_m": "2.4", "宽_m": "1.8",
        "周长_m": "8.4", "防水上翻高度_m": "0.3", "附加层面积_m2": "2.1",
        "涂膜厚度_mm": "1.5", "附加层宽度_m": "0.25", "搭接宽度_mm": "100",
        "卷材幅宽_mm": "1000", "间数": "1"
    }
    details, result = m.calculate_row(
        row, rules["做法"], rules["默认损耗率"], 0, project, standards
    )
    assert result["状态"] == "完成"
    assert result["来源"] == "项目规则/节点规则/JC-01"
    assert result["做法名称"] == "卫生间聚氨酯防水涂料（示例节点）"


def test_dedicated_missing_project_rule_does_not_fallback():
    row = {
        "部位类型": "防水", "做法编号": "F01", "做法类型": "项目专属", "节点编号": "JC-99",
        "节点详图": "节点图", "长_m": "2.4", "宽_m": "1.8", "间数": "1"
    }
    details, result = m.calculate_row(
        row, rules["做法"], rules["默认损耗率"], 0, project, standards
    )
    assert result["状态"] == "缺项"
    assert "项目专属做法未匹配节点规则 JC-99" in result["缺项说明"]


def test_duplicate_rows_flagged():
    row = {
        "楼层": "二层", "区域": "A区", "房间": "办公室", "部位类型": "地面/楼面",
        "做法编号": "D01", "做法名称": "地砖楼面", "做法类型": "通用",
        "长_m": "6", "宽_m": "4.5", "间数": "1", "依据": "建施-05"
    }
    issues = m.validate_rows([row, row])
    assert "与第 2 行完全重复，已跳过计算" in issues[1]


def test_invalid_numeric_and_required_fields():
    row = {
        "楼层": "二层", "区域": "A区", "房间": "办公室", "部位类型": "地面/楼面",
        "做法编号": "D01", "做法名称": "地砖楼面", "做法类型": "通用",
        "长_m": "6米", "间数": "0", "依据": "建施-05"
    }
    issues = m.validate_rows([row])[0]
    joined = "、".join(issues)
    assert any("长_m必须是数字" in x for x in issues), joined
    assert any("间数必须大于0" in x for x in issues), joined
    assert "宽_m" not in joined, joined


def test_excel_output_has_six_sheets():
    headers = [
        "楼层", "区域", "房间", "部位类型", "做法编号", "做法名称", "做法类型", "依据", "间数",
        "长_m", "宽_m"
    ]
    row = {
        "楼层": "二层", "区域": "A区", "房间": "办公室", "部位类型": "地面/楼面",
        "做法编号": "D01", "做法名称": "地砖楼面", "做法类型": "通用", "依据": "建施-05",
        "间数": "1", "长_m": "6", "宽_m": "4.5"
    }
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        input_path = td_path / "input.csv"
        out_dir = td_path / "out"
        m.write_csv(input_path, headers, [row])
        old_argv = sys.argv
        sys.argv = [
            "decor_qty.py", "--input", str(input_path), "--out-dir", str(out_dir),
            "--project", "测试项目"
        ]
        try:
            code = m.main()
        finally:
            sys.argv = old_argv
        assert code == 0
        from openpyxl import load_workbook
        wb = load_workbook(out_dir / "装饰概算物资计划台账.xlsx", read_only=True)
        assert wb.sheetnames == ["概算", "物资计划", "明细", "汇总", "缺项", "依据"]
        wb.close()


def test_quick_room_template_expands_practices():
    source = {
        "楼层": "二层", "区域": "A区", "房间": "办公室", "房间类型": "办公室",
        "长_m": "6", "宽_m": "4.5", "高_m": "3.1", "周长_m": "21", "间数": "1"
    }
    rows = m.expand_quick_rows([source], room_rules)
    assert len(rows) == 4
    assert [r["做法编号"] for r in rows] == ["D01", "W01", "C01", "T01"]
    assert rows[0]["源行号"] == "2"


def test_quick_unknown_room_becomes_issue():
    source = {"楼层": "二层", "区域": "A区", "房间": "包厢", "房间类型": "未知房间", "间数": "1"}
    rows = m.expand_quick_rows([source], room_rules)
    assert len(rows) == 1
    assert rows[0]["做法编号"] == ""
    assert "不在标准房间模板" in rows[0]["说明"]
    issues = m.validate_rows(rows)[0]
    assert any("做法编号为必填项" in x for x in issues)


def test_quick_input_generates_multiple_practices():
    source = {
        "楼层": "二层", "区域": "A区", "房间": "办公室", "房间类型": "办公室",
        "长_m": "6", "宽_m": "4.5", "高_m": "3.1", "周长_m": "21",
        "门窗洞口面积_m2": "2.1", "扣减长度_m": "0.9", "间数": "1", "依据": "建施-05"
    }
    rows = m.expand_quick_rows([source], room_rules)
    all_details = []
    for idx, row in enumerate(rows):
        details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], idx, standards=standards)
        all_details.extend(details)
        assert result["状态"] == "完成"
    assert len(all_details) >= 8


def test_estimate_uses_base_quantity_without_loss_overlap_additional():
    row = {
        "部位类型": "屋面防水", "做法编号": "F02", "长_m": "5", "宽_m": "6",
        "周长_m": "22", "防水上翻高度_m": "0.25", "附加层长度_m": "12",
        "附加层宽度_m": "0.5", "搭接宽度_mm": "100", "卷材幅宽_mm": "1000",
        "防水道数": "2", "涂膜厚度_mm": "1.5", "间数": "1", "依据": "建施-18"
    }
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    estimates = m.build_estimates([result], quantity_rules, price_rules)
    assert len(estimates) == 1
    assert abs(estimates[0]["清单工程量"] - (5 * 6 + 22 * 0.25)) < 1e-9
    membrane = next(d for d in details if d["材料计算方式"] == "overlap_area")
    assert membrane["订货量"] > estimates[0]["清单工程量"]


def test_estimate_amount_equals_quantity_times_price():
    row = {"部位类型": "地面/楼面", "做法编号": "D01", "长_m": "6", "宽_m": "4.5", "间数": "1"}
    _, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    estimate = m.build_estimates([result], quantity_rules, price_rules)[0]
    assert abs(estimate["清单工程量"] - 27) < 1e-9
    assert abs(estimate["概算合价"] - 27 * 180) < 1e-9


def test_material_plan_keeps_loss_demand():
    row = {"部位类型": "地面/楼面", "做法编号": "D01", "长_m": "6", "宽_m": "4.5", "间数": "1"}
    details, _ = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    plan = m.build_material_plan(details, material_rules)
    tile = next(r for r in plan if r["材料名称"] == "地砖")
    assert abs(tile["需求量"] - 27 * 1.03) < 1e-9
    assert abs(tile["净量"] - 27) < 1e-9
    assert tile["需求量"] > tile["净量"]


def test_membrane_difference_between_estimate_and_material_plan():
    row = {
        "部位类型": "屋面防水", "做法编号": "F02", "长_m": "5", "宽_m": "6",
        "周长_m": "22", "防水上翻高度_m": "0.25", "附加层长度_m": "12",
        "附加层宽度_m": "0.5", "搭接宽度_mm": "100", "卷材幅宽_mm": "1000",
        "防水道数": "2", "涂膜厚度_mm": "1.5", "间数": "1", "依据": "建施-18"
    }
    details, result = m.calculate_row(row, rules["做法"], rules["默认损耗率"], 0, standards=standards)
    estimate = m.build_estimates([result], quantity_rules, price_rules)[0]
    plan = m.build_material_plan(details, material_rules)
    membrane = next(r for r in plan if r["材料名称"] == "SBS改性沥青防水卷材")
    expected = estimate["清单工程量"] * (1000 / 900) * 1.1
    assert abs(membrane["需求量"] - expected) < 1e-9
    assert membrane["需求量"] > estimate["清单工程量"]


def test_missing_quantity_and_price_rules_are_flagged():
    result = {"行号": 2, "部位类型": "地面/楼面", "做法编号": "D01", "做法名称": "地砖楼面", "单位": "m2", "基面量": 27, "状态": "完成"}
    estimate = m.build_estimates([result], {}, {})[0]
    assert estimate["清单工程量"] == 0
    assert estimate["概算合价"] == 0
    assert "清单工程量规则" in estimate["缺项说明"]
    assert "综合单价" in estimate["缺项说明"]


def test_unmatched_material_classification_is_explicit():
    detail = {"材料名称": "未知板", "规格": "20mm", "材料单位": "m2", "损耗率": "0.05", "订货量": "10.5", "状态": "完成"}
    plan = m.build_material_plan([detail], {})
    assert plan[0]["材料类别"] == "其他"
    assert plan[0]["供货方式"] == "按项目确认"
    assert "物资分类规则未匹配" in plan[0]["备注"]



def test_standard_pack_has_fifteen_standards_and_metadata():
    assert standards["version"] == "0.6.0"
    assert standards["知识库来源"].endswith("00_装饰装修造价标准包索引.md")
    assert standards["引用边界"]
    assert len(standards["标准清单"]) == 15


def test_standard_pack_contains_energy_and_residential_standards():
    codes = {x["编号"] for x in standards["标准清单"]}
    assert {"GB 50411-2019", "GB 50327-2001"} <= codes


def test_gb50854_covers_core_decor_parts():
    item = next(x for x in standards["标准清单"] if x["编号"] == "GB 50854-2013")
    assert {"地面", "墙面", "顶棚", "踢脚", "防水", "保温"} <= set(item["适用部位"])


def test_all_quantity_rules_have_structured_standard_fields():
    required = {
        "计量单位", "计算规则", "依据", "技术规范", "构造要求",
        "复核点", "是否计入损耗"
    }
    assert set(quantity_rules["清单规则"]) == {"D01", "D02", "W01", "W02", "C01", "T01", "F01", "F02", "E01"}
    for code, rule in quantity_rules["清单规则"].items():
        assert required <= set(rule), (code, required - set(rule))
        assert rule["是否计入损耗"] is False
        assert rule["技术规范"] and rule["构造要求"] and rule["复核点"]


def test_quantity_units_match_practice_base_units():
    for code, rule in quantity_rules["清单规则"].items():
        assert rules["做法"][code]["基面"] == rule["计量单位"], code


def test_tile_floor_missing_binder_is_flagged():
    finish = {"名称": "地砖楼面", "材料": [{"名称": "地砖"}, {"名称": "填缝剂"}]}
    issues, checks = m._standard_checks({"部位类型": "地面/楼面"}, finish, standards)
    assert any("未列结合层/粘结材料" in x for x in issues)
    assert checks


def test_insulation_missing_anchor_or_mesh_is_flagged():
    finish = {"名称": "外墙外保温", "材料": [
        {"名称": "保温板"}, {"名称": "粘结砂浆"}, {"名称": "抹面砂浆"}
    ]}
    issues, checks = m._standard_checks({"部位类型": "外墙保温"}, finish, standards)
    assert any("未列锚栓/锚固件" in x for x in issues)
    assert any("未列网格布" in x for x in issues)
    assert checks


def test_paint_missing_primer_or_topcoat_is_flagged():
    finish = {"名称": "乳胶漆墙面", "材料": [{"名称": "抹灰砂浆"}, {"名称": "腻子"}]}
    issues, checks = m._standard_checks({"部位类型": "墙面"}, finish, standards)
    assert any("未列底漆" in x for x in issues)
    assert any("未列面漆" in x for x in issues)
    assert checks


def test_basis_includes_knowledge_index_and_all_standards():
    results = [{"行号": 2, "部位类型": "地面/楼面", "做法编号": "D01", "做法名称": "地砖楼面", "状态": "完成", "来源": "rules/装饰做法.json"}]
    basis = m._build_basis(results, standards)
    assert any(x["类型"] == "知识库索引" and x["编号/来源"] == standards["知识库来源"] for x in basis)
    standard_codes = {x["编号/来源"] for x in basis if x["类型"] == "国家规范参考"}
    assert len(standard_codes) == 15
    assert "GB 50411-2019" in standard_codes and "GB 50327-2001" in standard_codes


if __name__ == "__main__":
    failures = 0
    tests = [(name, fn) for name, fn in list(globals().items()) if name.startswith("test_") and callable(fn)]
    for name, fn in tests:
        try:
            fn()
            print(f"PASS {name}")
        except Exception as exc:
            failures += 1
            print(f"FAIL {name}: {exc}")
    if failures:
        raise SystemExit(f"{failures} test(s) failed")
    print(f"PASS all {len(tests)} tests")
