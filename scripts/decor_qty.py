#!/usr/bin/env python3
"""装饰工程快速概算与物资计划算量。

输入结构化明细表或快速房间表，输出概算、物资计划、Excel 台账、材料订货量、校验清单和复核报告。
概算采用清单/商务口径；物资计划采用采购口径。防水附加层、搭接、涂膜厚度等按内置国标参考要求校验；项目节点规则优先。
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

NUMBER_RE = re.compile(r"^[-+]?(\d+(\.\d*)?|\.\d+)$")
SPEC_RE = re.compile(r"(\d+(?:\.\d+)?)\s*[x×*]\s*(\d+(?:\.\d+)?)")

REQUIRED_COLUMNS = [
    "楼层", "区域", "房间", "部位类型", "做法编号", "做法名称", "做法类型", "依据", "间数"
]

NUMERIC_COLUMNS = [
    "长_m", "宽_m", "高_m", "周长_m", "门窗洞口面积_m2", "洞口侧壁增加_m2",
    "扣减长度_m", "直接面积_m2", "直接长度_m", "防水上翻高度_m",
    "附加层面积_m2", "附加层长度_m", "附加层宽度_m", "搭接宽度_mm",
    "卷材幅宽_mm", "涂膜厚度_mm", "防水道数", "间数"
]

QUICK_CARRY_FIELDS = [
    "楼层", "区域", "房间", "房间类型", "长_m", "宽_m", "高_m", "周长_m",
    "门窗洞口面积_m2", "洞口侧壁增加_m2", "扣减长度_m", "防水上翻高度_m",
    "附加层面积_m2", "附加层长度_m", "附加层宽度_m", "搭接宽度_mm",
    "卷材幅宽_mm", "涂膜厚度_mm", "防水道数", "间数", "说明", "依据", "需核对"
]


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\u3000", " ").strip()


def _num(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        v = float(value)
        return v if math.isfinite(v) else None
    s = _text(value).replace(",", "")
    if not s:
        return None
    if not NUMBER_RE.fullmatch(s):
        return None
    try:
        v = float(s)
        return v if math.isfinite(v) else None
    except ValueError:
        return None


def _num_or_zero(value: Any) -> float:
    v = _num(value)
    return v if v is not None else 0.0


def _num_default(value: Any, default: float) -> float:
    v = _num(value)
    return default if v is None else v


def _fmt(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return f"{value:.3f}".rstrip("0").rstrip(".")


def _parse_piece_area(spec: str) -> float | None:
    m = SPEC_RE.search(_text(spec))
    if not m:
        return None
    try:
        # 输入规格默认毫米，如 600x600。
        return (float(m.group(1)) / 1000.0) * (float(m.group(2)) / 1000.0)
    except ValueError:
        return None


def read_table(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 .xlsx 需要 openpyxl；请改用 CSV 或安装 openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [_text(v) for v in rows[0]]
        return [
            {h: _text(v) for h, v in zip(headers, row) if h}
            for row in rows[1:]
            if any(_text(v) for v in row)
        ]

    if suffix not in {".csv", ".tsv", ".txt"}:
        raise RuntimeError(f"暂不支持的输入类型：{suffix}（支持 CSV/TSV/XLSX）")
    delimiter = "\t" if suffix == ".tsv" else ","
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with path.open("r", encoding=encoding, newline="") as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                if not reader.fieldnames:
                    return []
                return [
                    {_text(k): _text(v) for k, v in row.items() if _text(k)}
                    for row in reader
                ]
        except UnicodeDecodeError:
            continue
    raise RuntimeError("CSV 编码无法识别；请另存为 UTF-8")


def write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _quick_issue_row(source: dict[str, str], room_type: str, message: str) -> dict[str, str]:
    row = {k: _text(source.get(k)) for k in QUICK_CARRY_FIELDS}
    row.update({
        "部位类型": "房间模板",
        "做法编号": "",
        "做法名称": "房间模板缺项",
        "做法类型": "通用",
        "说明": message,
        "房间类型": room_type,
    })
    return row


def is_quick_table(rows: list[dict[str, str]]) -> bool:
    fields: set[str] = set()
    for row in rows:
        fields.update(row.keys())
    return "房间类型" in fields and "做法编号" not in fields


def expand_quick_rows(
    rows: list[dict[str, str]],
    template_data: dict[str, Any] | None,
) -> list[dict[str, str]]:
    """把快速概算输入按房间类型展开为标准明细输入。"""
    templates = template_data or {}
    if "房间模板" in templates and isinstance(templates["房间模板"], dict):
        templates = templates["房间模板"]
    expanded: list[dict[str, str]] = []
    for idx, source in enumerate(rows):
        room_type = _text(source.get("房间类型"))
        template = templates.get(room_type)
        if not isinstance(template, dict):
            issue_row = _quick_issue_row(
                source, room_type,
                f"房间类型“{room_type or '(空)'}”不在标准房间模板中，请确认后改用明细输入"
            )
            issue_row["源行号"] = str(idx + 2)
            expanded.append(issue_row)
            continue

        practices = template.get("做法", [])
        if not isinstance(practices, list) or not practices:
            issue_row = _quick_issue_row(source, room_type, "房间模板未配置做法")
            issue_row["源行号"] = str(idx + 2)
            expanded.append(issue_row)
            continue

        for seq, practice in enumerate(practices, 1):
            if not isinstance(practice, dict):
                issue_row = _quick_issue_row(
                    source, room_type, f"第{seq}条做法配置不是对象"
                )
                issue_row["源行号"] = str(idx + 2)
                expanded.append(issue_row)
                continue
            new_row = {k: _text(source.get(k)) for k in QUICK_CARRY_FIELDS}
            for key in ("部位类型", "做法编号", "做法名称", "做法类型"):
                new_row[key] = _text(practice.get(key))
            defaults = practice.get("默认字段", {})
            if isinstance(defaults, dict):
                for key, value in defaults.items():
                    if not _text(new_row.get(key)):
                        new_row[key] = _text(value)
            if not new_row.get("做法类型"):
                new_row["做法类型"] = "通用"
            new_row["房间类型"] = room_type
            new_row["源行号"] = str(idx + 2)
            template_name = _text(template.get("名称")) or room_type
            old_note = new_row.get("说明", "")
            new_row["说明"] = f"房间模板展开：{template_name} / 第{seq}条做法" + (
                f"；{old_note}" if old_note else ""
            )
            expanded.append(new_row)
    return expanded


def validate_headers(rows: list[dict[str, str]]) -> list[str]:
    fields = set()
    for row in rows:
        fields.update(row.keys())
    missing = [c for c in REQUIRED_COLUMNS if c not in fields]
    if missing:
        return [f"缺少必填列：{ '、'.join(missing) }"]
    return []


def validate_rows(rows: list[dict[str, str]]) -> dict[int, list[str]]:
    issues: dict[int, list[str]] = defaultdict(list)
    signature_first: dict[str, int] = {}
    required_values = ["楼层", "区域", "房间", "部位类型", "做法编号", "做法名称", "做法类型", "依据", "间数"]

    for idx, row in enumerate(rows):
        line = int(_num(row.get("源行号")) or idx + 2)
        signature = "|".join(_text(row.get(k)) for k in sorted(row.keys()))
        if signature in signature_first:
            issues[idx].append(f"与第 {signature_first[signature]} 行完全重复，已跳过计算")
        else:
            signature_first[signature] = line

        for col in required_values:
            if not _text(row.get(col)):
                issues[idx].append(f"{col}为必填项")

        practice_type = _text(row.get("做法类型"))
        if practice_type and practice_type not in {"通用", "项目专属"}:
            issues[idx].append("做法类型只允许：通用 / 项目专属")
        if "专属" in practice_type:
            if not _text(row.get("节点编号")):
                issues[idx].append("项目专属做法缺少节点编号")
            if not _text(row.get("节点详图")):
                issues[idx].append("项目专属做法缺少节点详图")

        check = _text(row.get("需核对"))
        if check and check.lower() not in {"1", "y", "yes", "是", "否", "需核对", "true", "false"}:
            issues[idx].append("需核对只允许：是 / 否 / 需核对 / 空")

        for col in NUMERIC_COLUMNS:
            if col not in row:
                continue
            raw = _text(row.get(col))
            if not raw:
                continue
            value = _num(raw)
            if value is None:
                issues[idx].append(f"{col}必须是数字；列名已含单位，请勿填写文字单位")
            elif value < 0:
                issues[idx].append(f"{col}不能小于0")
            elif col == "间数" and value <= 0:
                issues[idx].append("间数必须大于0")

        part = _text(row.get("部位类型"))
        room = _text(row.get("房间"))
        if "防水" in part:
            if "附加层宽度_m" not in row:
                issues[idx].append("防水行缺少附加层宽度_m")
            if "搭接宽度_mm" not in row:
                issues[idx].append("防水行缺少搭接宽度_mm")
            if "涂膜厚度_mm" not in row:
                issues[idx].append("防水行缺少涂膜厚度_mm")
            if any(k in part + room for k in ("卫生间", "淋浴", "厨房")):
                upturn = _num(row.get("防水上翻高度_m"))
                if upturn is None:
                    issues[idx].append("卫生间/厨房防水缺少防水上翻高度_m")
                elif upturn <= 0:
                    issues[idx].append("卫生间/厨房防水上翻高度应大于0")
    return issues


def base_quantity(row: dict[str, str]) -> tuple[float, str, str, list[str], list[str]]:
    """返回基面量、单位、计算式、状态、缺项说明。"""
    part = _text(row.get("部位类型"))
    p = re.sub(r"[\s/／]", "", part)
    count = _num_or_zero(row.get("间数")) or 1.0
    length = _num(row.get("长_m"))
    width = _num(row.get("宽_m"))
    height = _num(row.get("高_m"))
    perimeter = _num(row.get("周长_m"))
    openings = _num_or_zero(row.get("门窗洞口面积_m2"))
    side_area = _num_or_zero(row.get("洞口侧壁增加_m2"))
    deduction_length = _num_or_zero(row.get("扣减长度_m"))
    direct_area = _num(row.get("直接面积_m2"))
    direct_length = _num(row.get("直接长度_m"))
    upturn = _num_or_zero(row.get("防水上翻高度_m"))
    missing: list[str] = []
    status = "完成"

    def miss(*fields: str) -> None:
        missing.extend(fields)
        nonlocal status
        status = "缺项"

    if direct_area is not None and direct_area > 0:
        area = direct_area * count
        formula = f"直接面积 {_fmt(direct_area)} × 间数 {_fmt(count)}"
        return area, "m2", formula, status, missing

    if direct_length is not None and direct_length > 0:
        length_q = direct_length * count
        formula = f"直接长度 {_fmt(direct_length)} × 间数 {_fmt(count)}"
        return length_q, "m", formula, status, missing

    if "踢脚" in p or "压顶" in p:
        if perimeter is not None:
            length_q = (perimeter - deduction_length) * count
            formula = f"({_fmt(perimeter)} - 扣减 {_fmt(deduction_length)}) × {_fmt(count)}"
            return max(length_q, 0.0), "m", formula, status, missing
        if length is not None and width is not None:
            length_q = (2 * (length + width) - deduction_length) * count
            formula = f"(2×({_fmt(length)}+{_fmt(width)}) - 扣减 {_fmt(deduction_length)}) × {_fmt(count)}"
            return max(length_q, 0.0), "m", formula, status, missing
        miss("周长_m 或 长_m+宽_m")
        return 0.0, "m", "", status, missing

    if "防水" in p:
        if length is not None and width is not None:
            flat = length * width
            if upturn > 0:
                if perimeter is None:
                    miss("防水上翻时的周长_m")
                    return flat * count, "m2", f"{_fmt(length)} × {_fmt(width)}；上翻缺周长", status, missing
                area = (flat + perimeter * upturn) * count
                formula = f"({_fmt(length)}×{_fmt(width)} + {_fmt(perimeter)}×上翻{_fmt(upturn)}) × {_fmt(count)}"
            else:
                area = flat * count
                formula = f"{_fmt(length)} × {_fmt(width)} × {_fmt(count)}"
            return max(area, 0.0), "m2", formula, status, missing
        miss("长_m+宽_m 或 直接面积_m2")
        return 0.0, "m2", "", status, missing

    if "地面" in p or "楼面" in p or "顶棚" in p or "吊顶" in p or "屋面" in p or "楼梯" in p:
        if length is None or width is None:
            miss("长_m/宽_m 或 直接面积_m2")
            return 0.0, "m2", "", status, missing
        area = length * width * count
        formula = f"{_fmt(length)} × {_fmt(width)} × {_fmt(count)}"
        return area, "m2", formula, status, missing

    if any(k in p for k in ("墙", "保温", "幕墙", "立面")):
        if perimeter is not None and height is not None:
            gross = perimeter * height
            formula = f"({_fmt(perimeter)} × {_fmt(height)} - 洞口 {_fmt(openings)} + 侧壁 {_fmt(side_area)}) × {_fmt(count)}"
        elif length is not None and height is not None:
            gross = length * height
            formula = f"({_fmt(length)} × {_fmt(height)} - 洞口 {_fmt(openings)} + 侧壁 {_fmt(side_area)}) × {_fmt(count)}"
        elif length is not None and width is not None and height is not None:
            gross = 2 * (length + width) * height
            formula = f"(2×({_fmt(length)}+{_fmt(width)}) × {_fmt(height)} - 洞口 {_fmt(openings)} + 侧壁 {_fmt(side_area)}) × {_fmt(count)}"
            status = "需核对" if status == "完成" else status
            missing.append("矩形周长推定需核对")
        else:
            miss("周长_m+高_m / 长_m+高_m / 直接面积_m2")
            return 0.0, "m2", "", status, missing
        area = max(gross - openings + side_area, 0.0) * count
        return area, "m2", formula, status, missing

    if length is not None and width is not None:
        area = length * width * count
        formula = f"{_fmt(length)} × {_fmt(width)} × {_fmt(count)}"
        return area, "m2", formula, status, missing
    if length is not None and height is not None:
        area = length * height * count
        formula = f"{_fmt(length)} × {_fmt(height)} × {_fmt(count)}"
        return area, "m2", formula, status, missing
    miss("能确定基面量的尺寸或 直接面积_m2/直接长度_m")
    return 0.0, "", "", status, missing


def _get_rule_value(row: dict[str, str], finish: dict[str, Any], key: str) -> float | None:
    value = _num(row.get(key))
    if value is not None:
        return value
    req = finish.get("防水要求", {})
    if isinstance(req, dict):
        return _num(req.get(key))
    return None


def _standard_refs(part: str, standards: dict[str, Any]) -> str:
    refs: list[str] = []
    for item in standards.get("标准清单", []):
        applicable = item.get("适用部位", [])
        if not applicable or any(k in part for k in applicable):
            refs.append(f"{item.get('编号','')} {item.get('名称','')}")
    return "；".join(dict.fromkeys(x.strip() for x in refs if x.strip()))


def _standard_checks(
    row: dict[str, str],
    finish: dict[str, Any],
    standards: dict[str, Any],
) -> tuple[list[str], list[str]]:
    issues: list[str] = []
    checks: list[str] = []
    part = _text(row.get("部位类型"))
    p = re.sub(r"[\s/／]", "", part)
    materials = finish.get("材料", []) if isinstance(finish.get("材料", []), list) else []
    material_names_list = [_text(x.get("名称")) for x in materials]
    material_names = "、".join(material_names_list)
    waterproof_section = standards.get("防水", {})
    min_lap = _num_default(waterproof_section.get("最低搭接宽度_mm"), 80.0)
    min_coating = _num_default(waterproof_section.get("最低涂膜厚度_mm"), 1.5)
    indoor_upturn = _num_default(waterproof_section.get("室内最低上翻高度_m"), 0.25)
    shower_upturn = _num_default(waterproof_section.get("淋浴区上翻高度_m"), 1.8)
    roof_min_count = _num_default(waterproof_section.get("屋面最低防水道数"), 2.0)

    if "防水" in p:
        has_membrane = any("卷材" in name for name in material_names_list)
        has_coating = any(("涂料" in name or "涂膜" in name) and "附加" not in name for name in material_names_list)
        has_additional = any("附加" in name for name in material_names_list)
        if not has_additional:
            issues.append("防水做法未列附加层材料")

        if has_membrane:
            lap = _get_rule_value(row, finish, "搭接宽度_mm")
            if lap is None:
                issues.append(f"防水卷材缺少搭接宽度_mm（内置参考下限 {min_lap}mm）")
            elif lap < min_lap:
                issues.append(f"防水卷材搭接宽度 {_fmt(lap)}mm 小于内置参考下限 {_fmt(min_lap)}mm")
            else:
                checks.append(f"搭接宽度 {lap}mm ≥ 参考下限 {min_lap}mm")
            roll_width = _get_rule_value(row, finish, "卷材幅宽_mm")
            if roll_width is None:
                issues.append("防水卷材缺少卷材幅宽_mm，无法计算搭接增量")
            elif roll_width <= 0:
                issues.append("卷材幅宽_mm必须大于0")

        if has_coating:
            thickness = _get_rule_value(row, finish, "涂膜厚度_mm")
            if thickness is None:
                issues.append(f"防水涂料缺少涂膜厚度_mm（内置参考下限 {min_coating}mm）")
            elif thickness < min_coating:
                issues.append(f"涂膜厚度 {_fmt(thickness)}mm 小于内置参考下限 {_fmt(min_coating)}mm")
            else:
                checks.append(f"涂膜厚度 {_fmt(thickness)}mm ≥ 参考下限 {_fmt(min_coating)}mm")

        width_limits = waterproof_section.get("部位最低附加层宽度_m", {})
        if not isinstance(width_limits, dict):
            width_limits = {}
        if "屋面" in p or "地下" in p:
            min_add_width = _num_default(width_limits.get("屋面/地下"), 0.5)
        else:
            min_add_width = _num_default(width_limits.get("室内"), 0.25)
        width = _get_rule_value(row, finish, "附加层宽度_m")
        if width is None:
            issues.append(f"防水附加层缺少附加层宽度_m（内置参考下限 {min_add_width}m）")
        elif width < min_add_width:
            issues.append(f"防水附加层宽度 {_fmt(width)}m 小于内置参考下限 {_fmt(min_add_width)}m")
        else:
            checks.append(f"附加层宽度 {_fmt(width)}m ≥ 参考下限 {_fmt(min_add_width)}m")

        upturn = _num(row.get("防水上翻高度_m"))
        room = _text(row.get("房间"))
        desc = part + room + _text(row.get("说明"))
        if any(k in desc for k in ("卫生间", "淋浴", "厨房")):
            min_upturn = shower_upturn if "淋浴" in desc else indoor_upturn
            if upturn is None or upturn <= 0:
                issues.append("室内有水房间防水上翻高度未填写或小于0")
            elif upturn < min_upturn:
                issues.append(f"室内有水房间防水上翻高度低于内置参考下限 {_fmt(min_upturn)}m；收头部位另按设计/规范核对")
            else:
                checks.append(f"防水上翻高度 {_fmt(upturn)}m ≥ 参考下限 {_fmt(min_upturn)}m")

        if "屋面" in p:
            count = _num(row.get("防水道数"))
            if count is None:
                issues.append("屋面防水缺少防水道数，无法核对设防要求")
            elif count < roof_min_count:
                issues.append(f"屋面防水道数少于内置参考值 {_fmt(roof_min_count)}道，应按防水等级和设计复核")
            else:
                checks.append(f"屋面防水道数 {_fmt(count)} 道 ≥ 参考值 {_fmt(roof_min_count)}道")

    elif any(k in p for k in ("地面", "楼面")):
        has_tile = any(k in material_names for k in ("地砖", "石材", "瓷砖", "块料"))
        if has_tile and not any(k in material_names for k in ("砂浆", "瓷砖胶", "粘结")):
            issues.append("块料地面未列结合层/粘结材料")
        if has_tile and "填缝" not in material_names:
            issues.append("块料地面未列填缝材料")
        checks.append("地面面层、结合层和块料规格按 GB 50209 与设计节点复核")
    elif "保温" in p:
        if not any(k in material_names for k in ("粘结", "砂浆")):
            issues.append("外墙保温未列粘结材料")
        if not any("锚栓" in name or "锚固" in name for name in material_names_list):
            issues.append("外墙保温未列锚栓/锚固件")
        if "网格布" not in material_names:
            issues.append("外墙保温未列网格布")
        if not any("抹面" in name for name in material_names_list):
            issues.append("外墙保温未列抹面砂浆")
        checks.append("外墙保温构造、锚固和防火隔离带按 JGJ 144、GB 50411 及设计文件复核")
    elif "墙" in p:
        has_tile = "砖" in material_names or "石材" in material_names
        if has_tile and not any(k in material_names for k in ("瓷砖胶", "砂浆", "粘结")):
            issues.append("块料墙面未列粘结材料")
        if has_tile and "填缝" not in material_names:
            issues.append("块料墙面未列填缝材料")
        paint_context = (
            any("乳胶漆" in name for name in material_names_list)
            or "乳胶漆" in _text(finish.get("名称"))
            or any("乳胶漆" in _text(row.get(k)) for k in ("做法名称", "说明", "做法类型"))
        )
        if paint_context:
            if not any("腻子" in name for name in material_names_list):
                issues.append("乳胶漆墙面未列腻子层")
            if not any("底漆" in name for name in material_names_list):
                issues.append("乳胶漆墙面未列底漆")
            if not any("面漆" in name for name in material_names_list):
                issues.append("乳胶漆墙面未列面漆")
        checks.append("墙面饰面粘结、空鼓和基层处理按 GB 50210 与设计节点复核")
    elif "吊顶" in p or "顶棚" in p:
        if "吊顶" in p:
            if not any(k in material_names for k in ("龙骨", "吊挂")):
                issues.append("吊顶做法未列龙骨/吊挂配件")
            if not any(k in material_names for k in ("板", "石膏板", "矿棉板", "铝板")):
                issues.append("吊顶做法未列饰面板")
        checks.append("吊顶龙骨、面板和检修口按 GB 50210 与节点详图复核")
    else:
        checks.append("装饰构造与材料按 GB 50210 及设计节点复核")
    return issues, checks


def _select_finish(
    row: dict[str, str],
    rules: dict[str, Any],
    project_data: dict[str, Any] | None,
) -> tuple[dict[str, Any] | None, str, list[str]]:
    code = _text(row.get("做法编号"))
    node_no = _text(row.get("节点编号"))
    practice_type = _text(row.get("做法类型"))
    project_data = project_data or {}
    issues: list[str] = []

    if "专属" in practice_type:
        node_rules = project_data.get("节点规则", {})
        if not isinstance(node_rules, dict) or node_no not in node_rules:
            issues.append(f"项目专属做法未匹配节点规则 {node_no or '(空)'}，不回退默认做法")
            return None, "", issues
        finish = dict(node_rules[node_no])
        rule_code = _text(finish.get("做法编号"))
        if rule_code and rule_code != code:
            issues.append(f"节点规则 {node_no} 对应做法编号 {rule_code}，与输入做法编号 {code} 不一致")
        return finish, f"项目规则/节点规则/{node_no}", issues

    default_finish = rules.get(code)
    if not default_finish:
        issues.append(f"做法编号 {code} 不在规则表")
        return None, "", issues
    override = project_data.get("做法覆盖", {}).get(code)
    if override:
        finish = dict(default_finish)
        finish.update(override)
        return finish, "项目规则/做法覆盖", issues
    return dict(default_finish), "rules/装饰做法.json", issues


def _additional_area(row: dict[str, str], finish: dict[str, Any], standards: dict[str, Any]) -> tuple[float | None, str, list[str]]:
    count = _num_or_zero(row.get("间数")) or 1.0
    direct = _num(row.get("附加层面积_m2"))
    if direct is not None and direct > 0:
        return direct * count, f"附加层面积 {_fmt(direct)} × 间数 {_fmt(count)}", []

    length = _num(row.get("附加层长度_m"))
    note_parts: list[str] = []
    if length is None:
        perimeter = _num(row.get("周长_m"))
        if perimeter is not None:
            length = perimeter
            note_parts.append("附加层长度默认按周长估算")
    width = _get_rule_value(row, finish, "附加层宽度_m")
    if width is None:
        width = _num_default(standards.get("防水", {}).get("默认附加层宽度_m"), 0.25)
        note_parts.append(f"附加层宽度默认 {_fmt(width)}m")
    if length is None or width is None or length <= 0 or width <= 0:
        return None, "", ["附加层缺少面积或长度/宽度"]
    formula = f"附加层长度 {_fmt(length)} × 附加层宽度 {_fmt(width)} × 间数 {_fmt(count)}"
    if note_parts:
        formula += "；" + "；".join(note_parts)
    return length * width * count, formula, note_parts

def calculate_row(
    row: dict[str, str],
    rules: dict[str, Any],
    default_losses: dict[str, float],
    row_index: int,
    project_data: dict[str, Any] | None = None,
    standards: dict[str, Any] | None = None,
    validation_issues: list[str] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    standards = standards or {}
    area, area_unit, formula, status, missing = base_quantity(row)
    code = _text(row.get("做法编号"))
    part = _text(row.get("部位类型"))
    practice_type = _text(row.get("做法类型")) or "通用"
    node_no = _text(row.get("节点编号"))
    node_drawing = _text(row.get("节点详图"))
    check = _text(row.get("需核对")).lower() in {"1", "y", "yes", "是", "需核对", "true"}
    if check:
        status = "需核对" if status == "完成" else status

    source_line = _num(row.get("源行号"))
    result: dict[str, Any] = {
        "行号": int(source_line) if source_line is not None else row_index + 2,
        "楼层": _text(row.get("楼层")),
        "区域": _text(row.get("区域")),
        "房间": _text(row.get("房间")),
        "房间类型": _text(row.get("房间类型")),
        "部位类型": part,
        "做法编号": code,
        "做法名称": _text(row.get("做法名称")),
        "做法类型": practice_type,
        "节点编号": node_no,
        "节点详图": node_drawing,
        "基面量": area,
        "单位": area_unit,
        "计算式": formula,
        "状态": status,
        "说明": _text(row.get("说明")),
        "依据": _text(row.get("依据")),
        "需核对": "是" if check else "",
        "缺项说明": "；".join(dict.fromkeys(missing + (validation_issues or []))),
        "规范依据": _standard_refs(part, standards),
        "防水校验": "",
        "来源": "rules/装饰做法.json",
    }
    details: list[dict[str, Any]] = []

    if validation_issues:
        result["状态"] = "重复" if any("重复" in x for x in validation_issues) else "缺项"
        return details, result

    if "专属" in practice_type and (not node_no or not node_drawing):
        result["状态"] = "缺项"
        result["缺项说明"] = "项目专属做法缺少节点编号或节点详图"
        return details, result

    finish, rules_source, select_issues = _select_finish(row, rules, project_data)
    if finish is None or select_issues:
        result["状态"] = "缺项"
        result["缺项说明"] = "；".join(dict.fromkeys([result["缺项说明"], *select_issues] if result["缺项说明"] else select_issues))
        result["来源"] = rules_source or "未匹配规则"
        return details, result

    if area <= 0:
        result["状态"] = "缺项"
        result["缺项说明"] = "；".join(dict.fromkeys([result["缺项说明"], "基面量不能计算"] if result["缺项说明"] else ["基面量不能计算"]))
        return details, result
    if not code:
        result["状态"] = "缺项"
        result["缺项说明"] = "；".join(dict.fromkeys([result["缺项说明"], "缺少做法编号"] if result["缺项说明"] else ["缺少做法编号"]))
        return details, result

    standard_issues, standard_checks = _standard_checks(row, finish, standards)
    if standard_issues:
        status = "缺项"
    result["状态"] = status
    result["缺项说明"] = "；".join(dict.fromkeys(([result["缺项说明"]] if result["缺项说明"] else []) + standard_issues))
    result["防水校验"] = "；".join(standard_checks)
    result["来源"] = rules_source
    result["做法名称"] = _text(finish.get("名称")) or _text(row.get("做法名称"))

    for item in finish.get("材料", []):
        name = _text(item.get("名称"))
        unit = _text(item.get("单位"))
        calc = _text(item.get("计算")).lower()
        factor = _num_or_zero(item.get("单耗")) or 1.0
        loss = _num_or_zero(item.get("损耗率"))
        if not loss:
            loss = _num_or_zero(default_losses.get(name))
        thickness = _num(item.get("厚度_mm"))
        spec = _text(item.get("规格"))
        piece_area = _num(item.get("单片面积_m2"))
        if piece_area is None or piece_area <= 0:
            piece_area = _parse_piece_area(spec)
        issue = ""
        formula2 = ""

        if not name or not unit:
            issue = "材料名称或单位缺失"
        elif calc in {"area", "m2"}:
            qty = area * factor * (1 + loss)
            formula2 = f"基面 {_fmt(area)} × 单耗 {_fmt(factor)} × (1+损耗 {_fmt(loss)})"
        elif calc in {"length", "m"}:
            if area_unit != "m":
                issue = "该材料按长度计算，但基面量不是长度"
                qty = 0.0
            else:
                qty = area * factor * (1 + loss)
                formula2 = f"基面长 {_fmt(area)} × 单耗 {_fmt(factor)} × (1+损耗 {_fmt(loss)})"
        elif calc in {"volume", "m3"}:
            if thickness is None or thickness <= 0:
                issue = "缺少厚度_mm"
                qty = 0.0
            else:
                qty = area * (thickness / 1000.0) * factor * (1 + loss)
                formula2 = f"基面 {_fmt(area)} × 厚 {_fmt(thickness)}/1000 × 单耗 {_fmt(factor)} × (1+损耗 {_fmt(loss)})"
        elif calc in {"pieces", "块"}:
            if piece_area is None or piece_area <= 0:
                issue = "缺少可解析的单片规格"
                qty = 0.0
            else:
                qty = math.ceil((area * factor * (1 + loss)) / piece_area)
                formula2 = f"向上取整(基面 {_fmt(area)} × 单耗 {_fmt(factor)} × (1+损耗 {_fmt(loss)}) ÷ 单片 {_fmt(piece_area)}㎡)"
        elif calc in {"count", "套"}:
            qty = area * factor * (1 + loss)
            formula2 = f"基面 {_fmt(area)} × 每单位数量 {_fmt(factor)} × (1+损耗 {_fmt(loss)})"
        elif calc in {"overlap_area", "搭接面积"}:
            lap = _get_rule_value(row, finish, "搭接宽度_mm")
            roll_width = _get_rule_value(row, finish, "卷材幅宽_mm")
            if lap is None or roll_width is None:
                issue = "缺少搭接宽度_mm或卷材幅宽_mm"
                qty = 0.0
            elif lap <= 0 or roll_width <= 0:
                issue = "搭接宽度_mm和卷材幅宽_mm必须大于0"
                qty = 0.0
            elif lap >= roll_width:
                issue = "搭接宽度_mm必须小于卷材幅宽_mm"
                qty = 0.0
            else:
                overlap_factor = roll_width / (roll_width - lap)
                qty = area * factor * overlap_factor * (1 + loss)
                formula2 = f"基面 {_fmt(area)} × 搭接系数 {_fmt(roll_width)}/({_fmt(roll_width)}-{_fmt(lap)}) × 单耗 {_fmt(factor)} × (1+损耗 {_fmt(loss)})"
        elif calc in {"additional_area", "附加面积"}:
            add_area, add_formula, add_notes = _additional_area(row, finish, standards)
            if add_area is None:
                issue = "；".join(add_notes or ["无法计算附加层面积"])
                qty = 0.0
            else:
                qty = add_area * factor * (1 + loss)
                formula2 = f"{add_formula} × 单耗 {_fmt(factor)} × (1+损耗 {_fmt(loss)})"
        else:
            issue = f"未知计算方式 {calc}"
            qty = 0.0

        if issue:
            result["状态"] = "缺项"
            old = result["缺项说明"].split("；") if result["缺项说明"] else []
            result["缺项说明"] = "；".join(dict.fromkeys(old + [issue]))
            continue
        detail = {
            **result,
            "材料名称": name,
            "规格": spec,
            "材料单位": unit,
            "材料计算方式": calc,
            "厚度_mm": thickness if thickness is not None else "",
            "损耗率": loss,
            "订货量": qty,
            "材料计算式": formula2,
        }
        details.append(detail)
    return details, result


def build_estimates(
    results: list[dict[str, Any]],
    quantity_rules: dict[str, Any],
    price_rules: dict[str, Any],
) -> list[dict[str, Any]]:
    """按清单/商务口径生成不含损耗、搭接和附加层的概算量。"""
    quantities = quantity_rules.get("清单规则", {}) if isinstance(quantity_rules, dict) else {}
    prices = price_rules.get("综合单价", {}) if isinstance(price_rules, dict) else {}
    rows: list[dict[str, Any]] = []
    for r in results:
        code = _text(r.get("做法编号"))
        q = quantities.get(code) if isinstance(quantities, dict) else None
        price = prices.get(code) if isinstance(prices, dict) else None
        q = q if isinstance(q, dict) else None
        price = price if isinstance(price, dict) else None
        issues = [_text(x) for x in (_text(r.get("缺项说明")).split("；") if _text(r.get("缺项说明")) else []) if x]
        status = _text(r.get("状态"))

        if q is None:
            issues.append(f"缺少做法编号 {code or '(空)'} 的清单工程量规则")
            status = "缺项" if status == "完成" else status
        q_unit = _text(q.get("计量单位")) if q else ""
        base_unit = _text(r.get("单位"))
        qty = _num_or_zero(r.get("基面量")) if q is not None and q_unit == base_unit else 0.0
        if q is not None and q_unit != base_unit:
            issues.append(f"清单单位 {q_unit or '(空)'} 与基面量单位 {base_unit or '(空)'} 不一致")
            status = "需核对" if status == "完成" else status

        price_unit = _text(price.get("单位")) if price else ""
        unit_price = _num(price.get("综合单价")) if price else None
        if price is None:
            issues.append(f"缺少做法编号 {code or '(空)'} 的综合单价")
            status = "需核对" if status == "完成" else status
        elif unit_price is None:
            issues.append("综合单价必须是数字")
            status = "需核对" if status == "完成" else status
            unit_price = 0.0
        elif price_unit and q_unit and price_unit != q_unit:
            issues.append(f"单价单位 {price_unit} 与清单单位 {q_unit} 不一致")
            status = "需核对" if status == "完成" else status
            unit_price = 0.0

        amount = qty * unit_price if q is not None and unit_price is not None else 0.0
        basis_parts = []
        if q:
            basis_parts.append(_text(q.get("依据")))
            for spec in q.get("技术规范", []) if isinstance(q.get("技术规范"), list) else []:
                basis_parts.append(_text(spec))
        basis_parts.append(_text(r.get("规范依据")))
        rows.append({
            "行号": r.get("行号"),
            "楼层": r.get("楼层"),
            "区域": r.get("区域"),
            "房间": r.get("房间"),
            "房间类型": r.get("房间类型"),
            "部位类型": r.get("部位类型"),
            "做法编号": code,
            "做法名称": r.get("做法名称"),
            "清单编码": _text(q.get("清单编码")) if q else "",
            "清单名称": _text(q.get("清单名称")) if q else "",
            "清单工程量": qty,
            "清单单位": q_unit,
            "综合单价": unit_price if unit_price is not None else "",
            "概算合价": amount,
            "状态": status,
            "单价来源": _text(price.get("单价来源")) if price else "",
            "依据": r.get("依据"),
            "规范依据": "；".join(dict.fromkeys([x for x in basis_parts if x])),
            "计算规则": _text(q.get("计算规则")) if q else "",
            "计算式": r.get("计算式"),
            "缺项说明": "；".join(dict.fromkeys(issues)),
        })
    return rows


def summarize_estimates(estimate_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    totals: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for r in estimate_rows:
        key = (
            _text(r.get("清单编码")),
            _text(r.get("清单名称")),
            _text(r.get("清单单位")),
            _text(r.get("单价来源")),
        )
        if key not in totals:
            totals[key] = {
                "清单编码": key[0],
                "清单名称": key[1],
                "清单单位": key[2],
                "单价来源": key[3],
                "清单工程量": 0.0,
                "概算合价": 0.0,
                "完成状态概算合价": 0.0,
                "涉及做法": set(),
            }
        qty = _num_or_zero(r.get("清单工程量"))
        amount = _num_or_zero(r.get("概算合价"))
        totals[key]["清单工程量"] += qty
        totals[key]["概算合价"] += amount
        if _text(r.get("状态")) == "完成":
            totals[key]["完成状态概算合价"] += amount
        if _text(r.get("做法编号")):
            totals[key]["涉及做法"].add(_text(r.get("做法编号")))
    out = []
    for v in totals.values():
        v["涉及做法"] = "、".join(sorted(v["涉及做法"]))
        out.append(v)
    return sorted(out, key=lambda x: (x["清单编码"], x["清单名称"]))


def build_material_plan(
    details: list[dict[str, Any]],
    material_rules: dict[str, Any],
) -> list[dict[str, Any]]:
    """按采购口径汇总材料需求量，保留损耗、搭接、附加层和块料取整结果。"""
    classifications = material_rules.get("物资分类", {}) if isinstance(material_rules, dict) else {}
    classifications = classifications if isinstance(classifications, dict) else {}
    totals: dict[tuple[str, str, str], dict[str, Any]] = {}
    for d in details:
        key = (_text(d.get("材料名称")), _text(d.get("规格")), _text(d.get("材料单位")))
        qty = _num_or_zero(d.get("订货量"))
        if key not in totals:
            totals[key] = {
                "净量": 0.0,
                "需求量": 0.0,
                "完成状态需求量": 0.0,
                "涉及部位": set(),
                "涉及楼层": set(),
            }
        loss_rate = _num_or_zero(d.get("损耗率"))
        totals[key]["净量"] += qty / (1 + loss_rate) if loss_rate >= 0 else qty
        totals[key]["需求量"] += qty
        if _text(d.get("状态")) == "完成":
            totals[key]["完成状态需求量"] += qty
        if _text(d.get("部位类型")):
            totals[key]["涉及部位"].add(_text(d.get("部位类型")))
        if _text(d.get("楼层")):
            totals[key]["涉及楼层"].add(_text(d.get("楼层")))

    category_order = {"主材": 0, "辅材": 1, "其他": 2}
    rows: list[dict[str, Any]] = []
    for (name, spec, unit), v in totals.items():
        cls = classifications.get(name) if isinstance(classifications, dict) else None
        cls = cls if isinstance(cls, dict) else {}
        demand = v["需求量"]
        net = v["净量"]
        loss_qty = max(demand - net, 0.0)
        loss_rate = loss_qty / net if net > 0 else 0.0
        category = _text(cls.get("材料类别")) or "其他"
        notes = _text(cls.get("备注"))
        if not cls:
            notes = "物资分类规则未匹配，请确认材料类别和供货方式"
        rows.append({
            "材料类别": category,
            "材料名称": name,
            "规格型号": spec,
            "单位": unit,
            "净量": net,
            "损耗率": loss_rate,
            "损耗量": loss_qty,
            "需求量": demand,
            "完成状态需求量": v["完成状态需求量"],
            "涉及部位": "、".join(sorted(v["涉及部位"])),
            "涉及楼层": "、".join(sorted(v["涉及楼层"])),
            "供货方式": _text(cls.get("供货方式")) or "按项目确认",
            "需求日期": "",
            "备注": notes,
        })
    return sorted(rows, key=lambda x: (category_order.get(x["材料类别"], 3), x["材料名称"], x["规格型号"], x["单位"]))


def aggregate_details(details: list[dict[str, Any]]) -> list[dict[str, Any]]:
    totals: dict[tuple[str, str, str], dict[str, Any]] = {}
    for d in details:
        key = (_text(d.get("材料名称")), _text(d.get("规格")), _text(d.get("材料单位")))
        if key not in totals:
            totals[key] = {
                "材料名称": key[0],
                "规格": key[1],
                "单位": key[2],
                "订货量": 0.0,
                "完成状态订货量": 0.0,
                "涉及部位": set(),
                "涉及楼层": set(),
            }
        qty = _num_or_zero(d.get("订货量"))
        totals[key]["订货量"] += qty
        if _text(d.get("状态")) == "完成":
            totals[key]["完成状态订货量"] += qty
        if _text(d.get("部位类型")):
            totals[key]["涉及部位"].add(_text(d.get("部位类型")))
        if _text(d.get("楼层")):
            totals[key]["涉及楼层"].add(_text(d.get("楼层")))
    out = []
    for v in totals.values():
        v["涉及部位"] = "、".join(sorted(v["涉及部位"]))
        v["涉及楼层"] = "、".join(sorted(v["涉及楼层"]))
        out.append(v)
    return sorted(out, key=lambda x: (x["材料名称"], x["规格"], x["单位"]))


def _write_excel(
    path: Path,
    estimates: list[dict[str, Any]],
    estimate_summary: list[dict[str, Any]],
    details: list[dict[str, Any]],
    summary: list[dict[str, Any]],
    missing: list[dict[str, Any]],
    basis: list[dict[str, Any]],
    material_plan: list[dict[str, Any]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def add_sheet(title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx, header in enumerate(headers, 1):
            max_len = len(header)
            for row in rows:
                max_len = max(max_len, len(_text(row.get(header))))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42)

    detail_headers = [
        "行号", "楼层", "区域", "房间", "部位类型", "做法编号", "做法名称", "做法类型",
        "节点编号", "节点详图", "基面量", "单位", "计算式", "材料名称", "规格", "材料单位",
        "材料计算方式", "厚度_mm", "损耗率", "订货量", "材料计算式", "状态", "缺项说明",
        "说明", "依据", "需核对", "规范依据", "防水校验", "来源"
    ]
    summary_headers = ["材料名称", "规格", "单位", "订货量", "完成状态订货量", "涉及部位", "涉及楼层"]
    missing_headers = [
        "行号", "楼层", "区域", "房间", "部位类型", "做法编号", "做法名称", "做法类型",
        "节点编号", "节点详图", "基面量", "单位", "计算式", "状态", "缺项说明", "说明",
        "依据", "需核对", "规范依据", "防水校验", "来源"
    ]
    estimate_headers = [
        "行号", "楼层", "区域", "房间", "房间类型", "部位类型", "做法编号", "做法名称",
        "清单编码", "清单名称", "清单工程量", "清单单位", "综合单价", "概算合价", "状态",
        "单价来源", "依据", "规范依据", "计算规则", "计算式", "缺项说明"
    ]
    estimate_summary_headers = [
        "清单编码", "清单名称", "清单单位", "单价来源", "清单工程量", "概算合价",
        "完成状态概算合价", "涉及做法"
    ]
    material_headers = [
        "材料类别", "材料名称", "规格型号", "单位", "净量", "损耗率", "损耗量", "需求量",
        "完成状态需求量", "涉及部位", "涉及楼层", "供货方式", "需求日期", "备注"
    ]
    basis_headers = ["类型", "关联行号", "编号/来源", "名称/内容", "关键要求", "适用部位", "算量作用", "状态"]
    add_sheet("概算", estimate_headers, estimates)
    add_sheet("物资计划", material_headers, material_plan)
    add_sheet("明细", detail_headers, details)
    add_sheet("汇总", summary_headers, summary)
    add_sheet("缺项", missing_headers, missing)
    add_sheet("依据", basis_headers, basis)
    wb.save(path)


def _build_basis(results: list[dict[str, Any]], standards: dict[str, Any]) -> list[dict[str, Any]]:
    basis: list[dict[str, Any]] = []
    knowledge_source = _text(standards.get("知识库来源"))
    if knowledge_source:
        basis.append({
            "类型": "知识库索引",
            "关联行号": "",
            "编号/来源": knowledge_source,
            "名称/内容": "装饰装修与造价标准包索引",
            "关键要求": "只引用标准元数据和算量约束，不复制规范原文。",
            "适用部位": "全部",
            "算量作用": "约束依据可追溯性",
            "状态": standards.get("更新日期", ""),
        })
    for r in results:
        source = _text(r.get("来源"))
        drawing = _text(r.get("节点详图"))
        basis.append({
            "类型": "项目依据",
            "关联行号": r.get("行号"),
            "编号/来源": source,
            "名称/内容": f"{r.get('部位类型')} / {r.get('做法编号')} {r.get('做法名称')}",
            "关键要求": f"输入依据：{r.get('依据')}；节点：{r.get('节点编号')} {drawing}",
            "适用部位": r.get("部位类型", ""),
            "算量作用": "基面量与材料配套复核",
            "状态": r.get("状态"),
        })
    for item in standards.get("标准清单", []):
        basis.append({
            "类型": "国家规范参考",
            "关联行号": "",
            "编号/来源": item.get("编号", ""),
            "名称/内容": item.get("名称", ""),
            "关键要求": item.get("关键要求", ""),
            "适用部位": "、".join(item.get("适用部位", [])) if isinstance(item.get("适用部位"), list) else item.get("适用部位", ""),
            "算量作用": item.get("算量作用", ""),
            "状态": "现行版本以项目/验收为准",
        })
    return basis


def main() -> int:
    base = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="装饰工程快速概算与物资计划算量")
    parser.add_argument("--input", required=True, help="输入 CSV/TSV/XLSX")
    parser.add_argument("--out-dir", required=True, help="输出目录")
    parser.add_argument("--rules", default=str(base / "rules" / "装饰做法.json"))
    parser.add_argument("--standards", default=str(base / "rules" / "国家规范要求.json"))
    parser.add_argument("--room-templates", default=str(base / "rules" / "标准房间模板.json"))
    parser.add_argument("--quantity-rules", default=str(base / "rules" / "清单工程量规则.json"))
    parser.add_argument("--price-rules", default=str(base / "rules" / "概算综合单价.json"))
    parser.add_argument("--material-rules", default=str(base / "rules" / "物资分类规则.json"))
    parser.add_argument("--project-rules", default="", help="项目专属规则 JSON，节点规则优先")
    parser.add_argument("--project", default="装饰工程")
    parser.add_argument("--level", default="", help="报表标题用楼层；如需筛选请先用表格筛选")
    parser.add_argument("--note", default="快速概算与物资计划口径；正式清单、计价和采购以合同清单、地方规则及设计文件为准。")
    args = parser.parse_args()

    out_dir = Path(args.out_dir).expanduser().resolve()
    paths = {
        "输入": Path(args.input).expanduser().resolve(),
        "材料规则": Path(args.rules).expanduser().resolve(),
        "规范规则": Path(args.standards).expanduser().resolve(),
        "房间模板": Path(args.room_templates).expanduser().resolve(),
        "清单规则": Path(args.quantity_rules).expanduser().resolve(),
        "单价规则": Path(args.price_rules).expanduser().resolve(),
        "物资分类": Path(args.material_rules).expanduser().resolve(),
    }
    if args.project_rules:
        paths["项目规则"] = Path(args.project_rules).expanduser().resolve()
    for label, path in paths.items():
        if not path.exists():
            print(f"❌文件不存在（{label}）：{path}", file=sys.stderr)
            return 2

    try:
        raw_rows = read_table(paths["输入"])
        with paths["材料规则"].open("r", encoding="utf-8") as f:
            rule_data = json.load(f)
        with paths["规范规则"].open("r", encoding="utf-8") as f:
            standards = json.load(f)
        with paths["房间模板"].open("r", encoding="utf-8") as f:
            room_templates = json.load(f)
        with paths["清单规则"].open("r", encoding="utf-8") as f:
            quantity_rules = json.load(f)
        with paths["单价规则"].open("r", encoding="utf-8") as f:
            price_rules = json.load(f)
        with paths["物资分类"].open("r", encoding="utf-8") as f:
            material_rules = json.load(f)
        project_data: dict[str, Any] = {}
        if "项目规则" in paths:
            with paths["项目规则"].open("r", encoding="utf-8") as f:
                project_data = json.load(f)
    except Exception as exc:
        print(f"❌读取失败：{exc}", file=sys.stderr)
        return 2

    if not raw_rows:
        print("❌输入表没有数据行", file=sys.stderr)
        return 2

    if is_quick_table(raw_rows):
        rows = expand_quick_rows(raw_rows, room_templates)
        input_mode = "快速房间输入"
    else:
        rows = raw_rows
        input_mode = "明细输入"

    header_issues = validate_headers(rows)
    if header_issues:
        print("❌表头校验失败：" + "；".join(header_issues), file=sys.stderr)
        return 2

    validation = validate_rows(rows)
    rules = rule_data.get("做法", {})
    default_losses = rule_data.get("默认损耗率", {})
    all_details: list[dict[str, Any]] = []
    results: list[dict[str, Any]] = []
    for i, row in enumerate(rows):
        details, result = calculate_row(
            row, rules, default_losses, i, project_data, standards, validation.get(i, [])
        )
        all_details.extend(details)
        results.append(result)

    estimates = build_estimates(results, quantity_rules, price_rules)
    estimate_summary = summarize_estimates(estimates)
    material_plan = build_material_plan(all_details, material_rules)
    missing = [r for r in results if r["状态"] != "完成"]
    summary = aggregate_details(all_details)
    basis = _build_basis(results, standards)

    out_dir.mkdir(parents=True, exist_ok=True)
    detail_headers = [
        "行号", "楼层", "区域", "房间", "部位类型", "做法编号", "做法名称",
        "做法类型", "节点编号", "节点详图", "基面量", "单位", "计算式", "材料名称", "规格",
        "材料单位", "材料计算方式", "厚度_mm", "损耗率", "订货量", "材料计算式", "状态",
        "缺项说明", "说明", "依据", "需核对", "规范依据", "防水校验", "来源"
    ]
    summary_headers = ["材料名称", "规格", "单位", "订货量", "完成状态订货量", "涉及部位", "涉及楼层"]
    missing_headers = [
        "行号", "楼层", "区域", "房间", "部位类型", "做法编号", "做法名称",
        "做法类型", "节点编号", "节点详图", "基面量", "单位", "计算式", "状态",
        "缺项说明", "说明", "依据", "需核对", "规范依据", "防水校验", "来源"
    ]
    estimate_headers = [
        "行号", "楼层", "区域", "房间", "房间类型", "部位类型", "做法编号", "做法名称",
        "清单编码", "清单名称", "清单工程量", "清单单位", "综合单价", "概算合价", "状态",
        "单价来源", "依据", "规范依据", "计算规则", "计算式", "缺项说明"
    ]
    estimate_summary_headers = [
        "清单编码", "清单名称", "清单单位", "单价来源", "清单工程量", "概算合价",
        "完成状态概算合价", "涉及做法"
    ]
    material_headers = [
        "材料类别", "材料名称", "规格型号", "单位", "净量", "损耗率", "损耗量", "需求量",
        "完成状态需求量", "涉及部位", "涉及楼层", "供货方式", "需求日期", "备注"
    ]
    basis_headers = ["类型", "关联行号", "编号/来源", "名称/内容", "关键要求", "适用部位", "算量作用", "状态"]
    write_csv(out_dir / "概算明细.csv", estimate_headers, estimates)
    write_csv(out_dir / "概算汇总.csv", estimate_summary_headers, estimate_summary)
    write_csv(out_dir / "物资计划.csv", material_headers, material_plan)
    write_csv(out_dir / "明细.csv", detail_headers, all_details)
    write_csv(out_dir / "材料汇总.csv", summary_headers, summary)
    write_csv(out_dir / "缺项核对.csv", missing_headers, missing)
    write_csv(out_dir / "依据.csv", basis_headers, basis)

    excel_path = out_dir / "装饰概算物资计划台账.xlsx"
    try:
        _write_excel(
            excel_path, estimates, estimate_summary, all_details, summary,
            missing, basis, material_plan
        )
    except Exception as exc:
        print(f"❌Excel台账生成失败：{exc}", file=sys.stderr)
        return 2

    report: list[str] = []
    report.append(f"# {args.project} · 装饰概算物资计划报告")
    report.append("")
    report.append("- 输入模式：" + input_mode)
    report.append(f"- 输入源数据行数：{len(raw_rows)}")
    report.append(f"- 展开后计算行数：{len(rows)}")
    report.append(f"- 明细材料行数：{len(all_details)}")
    report.append(f"- 概算行数：{len(estimates)}")
    report.append(f"- 物资计划项数：{len(material_plan)}")
    report.append(f"- 缺项/需核对/重复行数：{len(missing)}")
    if args.level:
        report.append(f"- 楼层：{args.level}")
    if args.note:
        report.append(f"- 说明：{args.note}")
    report.append("")
    report.append("## 口径说明")
    report.append("")
    if standards.get("知识库来源"):
        report.append(f"- 知识库索引：{standards.get('知识库来源')}（更新日期：{standards.get('更新日期', '')}）")
    report.append("- 清单工程量不包含损耗、卷材搭接和防水附加层；物资计划量包含采购增量。")
    report.append("- 项目合同清单、地方规则、现行标准、设计文件和节点详图优先于内置参考。")
    report.append("- 概算量采用清单/商务口径：取基面量，不含损耗、卷材搭接和防水附加层。")
    report.append("- 物资计划量采用采购口径：含损耗、卷材搭接、防水附加层，块料按单片面积向上取整。")
    report.append("- 清单编码和综合单价均为内置参考；合同清单、地方规则、企业定额、设计文件和现行规范优先。")
    report.append("- 本输出不用于竣工结算、对量或正式报价。")
    report.append("")
    report.append("## 概算汇总")
    report.append("")
    if estimate_summary:
        report.append("| 清单编码 | 清单名称 | 单位 | 清单工程量 | 概算合价 | 完成状态概算合价 | 涉及做法 |")
        report.append("|---|---|---|---:|---:|---:|---|")
        for r in estimate_summary:
            report.append(
                f"| {r['清单编码']} | {r['清单名称']} | {r['清单单位']} | {_fmt(r['清单工程量'])} | "
                f"{_fmt(r['概算合价'])} | {_fmt(r['完成状态概算合价'])} | {r['涉及做法']} |"
            )
    else:
        report.append("无可汇总概算项。")
    report.append("")
    report.append("## 物资计划汇总")
    report.append("")
    if material_plan:
        report.append("| 类别 | 材料名称 | 规格型号 | 单位 | 净量 | 损耗量 | 需求量 | 完成状态需求量 | 供货方式 |")
        report.append("|---|---|---|---|---:|---:|---:|---:|---|")
        for r in material_plan:
            report.append(
                f"| {r['材料类别']} | {r['材料名称']} | {r['规格型号']} | {r['单位']} | {_fmt(r['净量'])} | "
                f"{_fmt(r['损耗量'])} | {_fmt(r['需求量'])} | {_fmt(r['完成状态需求量'])} | {r['供货方式']} |"
            )
    else:
        report.append("无可汇总物资项。")
    report.append("")
    report.append("## 材料消耗汇总")
    report.append("")
    if summary:
        report.append("| 材料名称 | 规格 | 单位 | 订货量 | 完成状态订货量 | 涉及部位 | 涉及楼层 |")
        report.append("|---|---|---|---:|---:|---|---|")
        for s in summary:
            report.append(
                f"| {s['材料名称']} | {s['规格']} | {s['单位']} | {_fmt(s['订货量'])} | "
                f"{_fmt(s['完成状态订货量'])} | {s['涉及部位']} | {s['涉及楼层']} |"
            )
    else:
        report.append("无可汇总材料。")
    report.append("")
    report.append("## 待确认项")
    report.append("")
    if missing:
        report.append("| 行号 | 楼层 | 房间 | 部位 | 做法 | 状态 | 缺项说明 |")
        report.append("|---:|---|---|---|---|---|---|")
        for m in missing:
            report.append(
                f"| {m['行号']} | {m['楼层']} | {m['房间']} | {m['部位类型']} | "
                f"{m['做法编号']} {m['做法名称']} | {m['状态']} | {m['缺项说明']} |"
            )
    else:
        report.append("无。")
    report.append("")
    report.append("## 国家规范与计量计价参考")
    report.append("")
    for item in standards.get("标准清单", []):
        applicable = "、".join(item.get("适用部位", [])) if isinstance(item.get("适用部位"), list) else item.get("适用部位", "")
        report.append(f"- {item.get('编号','')} {item.get('名称','')}：{item.get('关键要求','')} 适用：{applicable}；算量作用：{item.get('算量作用','')}")
    report.append("")
    report.append("## 复核要求")
    report.append("")
    report.append("1. 核对楼层、区域、房间、房间类型、做法编号、节点编号和输入依据。")
    report.append("2. 核对门窗洞口、洞口侧壁、上翻高度、矩形周长推定和房间模板适用性。")
    report.append("3. 防水重点核对设防道数、材料类别、涂膜厚度、搭接、附加层、收头和节点。")
    report.append("4. 概算清单编码、项目特征和单价必须与合同清单、地方规则和项目商务口径核对。")
    report.append("5. 物资需求日期、批量、备货期、供货界面和企业物资目录按项目采购计划补充。")
    report.append("6. 项目专属做法必须使用节点规则；节点详图与默认规则冲突时，以节点详图为准。")
    report.append("")
    report.append("内置规范值和参考单价仅是算量提示，不代表规范条文或市场价原文；正式口径以现行规范、设计文件和项目约定为准。")
    (out_dir / "装饰概算物资计划报告.md").write_text("\n".join(report) + "\n", encoding="utf-8")

    print(f"✅已生成：{out_dir}")
    print(
        f"概算行：{len(estimates)}；物资计划项：{len(material_plan)}；"
        f"明细材料行：{len(all_details)}；缺项/需核对/重复行：{len(missing)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
